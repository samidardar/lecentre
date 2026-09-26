"""Extraction de texte brut : PDF, TXT/MD, DOCX, HTML, CSV (FAQ)."""
from __future__ import annotations

import csv
import html
import io
import re
from pathlib import PurePath

SUPPORTED = {"pdf", "txt", "md", "markdown", "docx", "html", "htm", "csv", "json"}


def source_type_for(filename: str) -> str:
    ext = PurePath(filename).suffix.lower().lstrip(".")
    if ext not in SUPPORTED:
        raise ValueError(f"format non supporté: .{ext} (supportés: {', '.join(sorted(SUPPORTED))})")
    return {"markdown": "md", "htm": "html"}.get(ext, ext)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            out.append(f"{'#' * int(level)} {text}")
        else:
            out.append(text)
    for table in d.tables:
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(out)


def _html(data: bytes) -> str:
    text = _decode(data)
    text = re.sub(r"(?is)<(script|style|noscript).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?i)<h([1-6])[^>]*>(.*?)</h\1>", lambda m: "\n\n" + "#" * int(m.group(1)) + " " + m.group(2) + "\n\n", text)
    text = re.sub(r"(?i)<(br|/p|/div|/li|/tr)[^>]*>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    return html.unescape(text)


def _csv(data: bytes) -> str:
    """CSV FAQ : colonnes question/réponse détectées, sinon une ligne = un enregistrement."""
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    if not rows:
        return ""
    cols = {c.lower().strip(): c for c in rows[0].keys() if c}
    q = next((cols[c] for c in cols if c in {"question", "q", "questions"}), None)
    a = next((cols[c] for c in cols if c in {"answer", "réponse", "reponse", "a", "response"}), None)
    if q and a:
        return "\n\n".join(f"Q: {r[q].strip()}\nR: {r[a].strip()}" for r in rows if r.get(q) and r.get(a))
    return "\n\n".join("; ".join(f"{k}: {v}" for k, v in r.items() if v) for r in rows)


def extract_text(filename: str, data: bytes) -> str:
    kind = source_type_for(filename)
    if kind == "pdf":
        raw = _pdf(data)
    elif kind == "docx":
        raw = _docx(data)
    elif kind == "html":
        raw = _html(data)
    elif kind == "csv":
        raw = _csv(data)
    else:
        raw = _decode(data)
    return clean_text(raw)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"-\n(?=\w)", "", text)  # césures PDF
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
